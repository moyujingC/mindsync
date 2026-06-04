#!/usr/bin/env python3
"""批量将 docx 文件转换为 markdown 格式"""

from docx import Document
from pathlib import Path
import os

def convert_docx_to_md(docx_path, output_dir=None):
    """将 docx 文件转换为 markdown"""
    doc = Document(docx_path)
    md_content = []

    # 处理段落
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        style_name = para.style.name if para.style else ''

        # 标题检测
        if 'Heading 1' in style_name:
            md_content.append(f'# {text}')
        elif 'Heading 2' in style_name:
            md_content.append(f'## {text}')
        elif 'Heading 3' in style_name:
            md_content.append(f'### {text}')
        elif 'Heading 4' in style_name:
            md_content.append(f'#### {text}')
        elif 'List' in style_name or 'Bullet' in style_name:
            md_content.append(f'- {text}')
        elif 'Number' in style_name:
            md_content.append(f'- {text}')
        else:
            md_content.append(text)

    # 处理表格
    for table in doc.tables:
        md_content.append('')
        header = table.rows[0]
        md_content.append('| ' + ' | '.join(cell.text.strip() for cell in header.cells) + ' |')
        md_content.append('|' + '|'.join('---' for _ in header.cells) + '|')
        for row in table.rows[1:]:
            md_content.append('| ' + ' | '.join(cell.text.strip() for cell in row.cells) + ' |')
        md_content.append('')

    # 生成输出路径
    docx_name = Path(docx_path).stem
    if output_dir:
        output_path = Path(output_dir) / f"{docx_name}.md"
    else:
        output_path = Path(docx_path).with_suffix('.md')

    # 写入文件
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write('\n\n'.join(md_content))

    return output_path

def main():
    source_dir = Path("/Users/xinran/Downloads/dev/mindsync/projects/aimandala/docs/sources/曼陀罗解读师【第3期】")
    output_dir = Path("/Users/xinran/Downloads/dev/mindsync/projects/aimandala/docs/sources/曼陀罗解读师【第3期】")

    docx_files = sorted(source_dir.glob("*.docx"))
    print(f"找到 {len(docx_files)} 个 docx 文件")

    for docx_file in docx_files:
        print(f"转换: {docx_file.name}")
        try:
            output_path = convert_docx_to_md(docx_file, output_dir)
            print(f"  -> {output_path.name}")
        except Exception as e:
            print(f"  -> 错误: {e}")

if __name__ == "__main__":
    main()