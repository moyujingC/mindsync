#!/usr/bin/env python3
import argparse
import csv
import json
import math
import os
import re
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


SAVE_URL = "https://openapi.biji.com/open/api/v1/resource/note/save"
PROGRESS_URL = "https://openapi.biji.com/open/api/v1/resource/note/task/progress"
DETAIL_URL = "https://openapi.biji.com/open/api/v1/resource/note/detail"
KB_ADD_URL = "https://openapi.biji.com/open/api/v1/resource/knowledge/note/batch-add"


def safe_slug(text: str) -> str:
    text = re.sub(r"[^\w\-\u4e00-\u9fff]+", "-", text, flags=re.UNICODE)
    text = re.sub(r"-+", "-", text).strip("-")
    return text or "note"


def get_headers():
    api_key = os.environ["GETNOTE_API_KEY"]
    client_id = os.environ["GETNOTE_CLIENT_ID"]
    return {
        "Authorization": api_key,
        "X-Client-ID": client_id,
        "Content-Type": "application/json",
    }


def request_json(url: str, method: str = "GET", data=None, timeout: int = 60):
    payload = None
    headers = get_headers()
    if data is not None:
        payload = json.dumps(data, ensure_ascii=False).encode("utf-8")
    req = urllib.request.Request(url, data=payload, headers=headers, method=method)
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        text = resp.read().decode("utf-8", errors="replace")
        return json.loads(text)


def read_rows(xlsx_path: Path, since: datetime):
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    idx = {name: i for i, name in enumerate(header)}
    selected = []
    for row in rows:
        published_at = row[idx["发布时间"]]
        if isinstance(published_at, datetime) and published_at >= since:
            selected.append(
                {
                    "note_id": str(row[idx["笔记ID"]]),
                    "note_link": (row[idx["笔记链接"]] or "").strip(),
                    "note_type": row[idx["笔记类型"]] or "",
                    "title": row[idx["笔记标题"]] or "",
                    "content": row[idx["笔记内容"]] or "",
                    "published_at": published_at.strftime("%Y-%m-%d %H:%M:%S"),
                    "video_duration": row[idx["笔记视频时长"]] or "",
                    "video_url": row[idx["笔记视频链接"]] or "",
                    "topics": row[idx["笔记话题"]] or "",
                }
            )
    selected.sort(key=lambda item: item["published_at"])
    return selected


def load_manifest(path: Path):
    if path.exists():
        return json.loads(path.read_text())
    return {}


def save_manifest(path: Path, manifest: dict):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2))


def submit_link_note(item: dict, retry_429_seconds: int = 70, max_attempts: int = 4):
    payload = {
        "title": item["title"],
        "note_type": "link",
        "link_url": item["note_link"],
        "tags": ["小红书", "亚慧AI产品经理", "2026", "应聘AI产品经理"],
        "parent_id": 0,
    }
    for attempt in range(max_attempts):
        try:
            body = request_json(SAVE_URL, method="POST", data=payload)
            if not body.get("success"):
                raise RuntimeError(json.dumps(body, ensure_ascii=False))
            tasks = ((body.get("data") or {}).get("tasks") or [])
            if not tasks:
                raise RuntimeError(f"save response missing tasks: {json.dumps(body, ensure_ascii=False)}")
            return tasks[0]["task_id"]
        except urllib.error.HTTPError as exc:
            if exc.code == 429 and attempt < max_attempts - 1:
                time.sleep(retry_429_seconds)
                continue
            raise RuntimeError(f"HTTP Error {exc.code}: {exc.reason}")
        except RuntimeError as exc:
            if "10202" in str(exc) and attempt < max_attempts - 1:
                time.sleep(retry_429_seconds)
                continue
            raise


def poll_task(task_id: str, sleep_seconds: int = 12, max_attempts: int = 150):
    for _ in range(max_attempts):
        body = request_json(PROGRESS_URL, method="POST", data={"task_id": task_id})
        if not body.get("success"):
            raise RuntimeError(json.dumps(body, ensure_ascii=False))
        data = body.get("data") or {}
        status = data.get("status")
        if status == "success":
            note_id = str(data.get("note_id") or "")
            if note_id and note_id != "0":
                return note_id
        if status == "failed":
            raise RuntimeError(data.get("error_msg") or f"task failed: {task_id}")
        time.sleep(sleep_seconds)
    raise RuntimeError(f"task timeout: {task_id}")


def fetch_detail(note_id: str):
    url = f"{DETAIL_URL}?{urllib.parse.urlencode({'id': note_id})}"
    body = request_json(url, method="GET")
    if not body.get("success"):
        raise RuntimeError(json.dumps(body, ensure_ascii=False))
    return (body.get("data") or {}).get("note") or {}


def write_markdown(md_path: Path, item: dict, detail: dict, getnote_note_id: str):
    md_path.parent.mkdir(parents=True, exist_ok=True)
    web_page = detail.get("web_page") or {}
    content = detail.get("content") or ""
    web_content = web_page.get("content") or ""
    excerpt = web_page.get("excerpt") or ""
    topics = detail.get("topics") or []
    tags = detail.get("tags") or []
    lines = [
        f"# {item['title']}",
        "",
        f"- 发布时间：`{item['published_at']}`",
        f"- 原小红书链接：{item['note_link']}",
        f"- Get笔记 Note ID：`{getnote_note_id}`",
        f"- 知识库：{', '.join(topic.get('name', '') for topic in topics if topic.get('name'))}",
        f"- 标签：{', '.join(tag.get('name', '') for tag in tags if tag.get('name'))}",
        "",
        "## Get笔记总结",
        "",
        content or "_无总结内容_",
        "",
        "## 原文摘录",
        "",
        excerpt or "_无摘录_",
        "",
        "## 抓取正文",
        "",
        web_content or "_无正文内容_",
        "",
    ]
    md_path.write_text("\n".join(lines))


def add_to_knowledge_base(topic_id: str, note_ids: list[str]):
    for i in range(0, len(note_ids), 20):
        chunk = note_ids[i : i + 20]
        body = request_json(KB_ADD_URL, method="POST", data={"topic_id": topic_id, "note_ids": chunk})
        if not body.get("success"):
            raise RuntimeError(json.dumps(body, ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--topic-id", required=True)
    parser.add_argument("--since", default="2026-01-01")
    parser.add_argument("--poll-seconds", type=int, default=12)
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--submit-interval-seconds", type=int, default=0)
    parser.add_argument("--retry-failed-only", action="store_true")
    parser.add_argument("--retry-429-seconds", type=int, default=70)
    parser.add_argument("--save-max-attempts", type=int, default=4)
    parser.add_argument("--poll-max-attempts", type=int, default=150)
    args = parser.parse_args()

    since = datetime.strptime(args.since, "%Y-%m-%d")
    xlsx_path = Path(args.xlsx).expanduser().resolve()
    base_dir = xlsx_path.parent
    export_dir = base_dir / "_markdown_from_getnote"
    meta_dir = base_dir / "_getnote_import_meta"
    manifest_path = meta_dir / "manifest.json"
    links_csv_path = meta_dir / "selected_rows.csv"

    rows = read_rows(xlsx_path, since)
    if args.limit > 0:
        rows = rows[: args.limit]

    meta_dir.mkdir(parents=True, exist_ok=True)
    with links_csv_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)

    manifest = load_manifest(manifest_path)

    created_note_ids = []
    for index, item in enumerate(rows, start=1):
        key = item["note_id"]
        entry = manifest.get(key, {})
        print(f"[{index}/{len(rows)}] {item['published_at']} {item['title']}")
        if args.retry_failed_only and entry and entry.get("status") not in {"failed", "submitted"}:
            continue
        if entry.get("status") == "done" and entry.get("getnote_note_id"):
            if not entry.get("knowledge_base_added"):
                add_to_knowledge_base(args.topic_id, [entry["getnote_note_id"]])
                entry["knowledge_base_added"] = True
                manifest[key] = entry
                save_manifest(manifest_path, manifest)
            created_note_ids.append(entry["getnote_note_id"])
            continue
        try:
            task_id = entry.get("task_id") or submit_link_note(
                item,
                retry_429_seconds=args.retry_429_seconds,
                max_attempts=args.save_max_attempts,
            )
            manifest[key] = {**entry, **item, "task_id": task_id, "status": "submitted", "updated_at": datetime.now().isoformat(timespec="seconds")}
            save_manifest(manifest_path, manifest)

            getnote_note_id = entry.get("getnote_note_id") or poll_task(
                task_id,
                sleep_seconds=args.poll_seconds,
                max_attempts=args.poll_max_attempts,
            )
            detail = fetch_detail(getnote_note_id)
            md_name = safe_slug(f"{item['published_at'][:10]}-{item['title']}") + ".md"
            md_path = export_dir / md_name
            write_markdown(md_path, item, detail, getnote_note_id)
            add_to_knowledge_base(args.topic_id, [getnote_note_id])

            manifest[key] = {
                **manifest[key],
                "status": "done",
                "getnote_note_id": getnote_note_id,
                "markdown_path": str(md_path),
                "knowledge_base_added": True,
                "updated_at": datetime.now().isoformat(timespec="seconds"),
            }
            save_manifest(manifest_path, manifest)
            created_note_ids.append(getnote_note_id)
            if args.submit_interval_seconds > 0:
                time.sleep(args.submit_interval_seconds)
        except Exception as exc:
            manifest[key] = {
                **entry,
                **item,
                "status": "failed",
                "error": str(exc),
                "updated_at": datetime.now().isoformat(timespec="seconds"),
            }
            save_manifest(manifest_path, manifest)
            print(f"  FAILED: {exc}", flush=True)

    unique_note_ids = []
    seen = set()
    for note_id in created_note_ids:
        if note_id and note_id not in seen:
            seen.add(note_id)
            unique_note_ids.append(note_id)

    add_to_knowledge_base(args.topic_id, unique_note_ids)
    print(f"Imported {len(unique_note_ids)} notes into topic {args.topic_id}")
    print(f"Manifest: {manifest_path}")
    print(f"Markdown dir: {export_dir}")


if __name__ == "__main__":
    main()
